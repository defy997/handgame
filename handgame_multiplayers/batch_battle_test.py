from __future__ import annotations

from dataclasses import asdict
from pathlib import Path
from typing import Optional

from openpyxl import Workbook

from heros.basic import BasicHero


ACTION_ORDER = {
    "self-attack": 0,
    "group-attack": 1,
    "anti-group-attack": 2,
    "attack": 3,
    "defense": 4,
    "control": 5,
    None: 6,
}

OUTPUT_FILE = Path(__file__).with_name("battle_results.xlsx")


def action_name(hero: BasicHero, action_id: int) -> str:
    return next(name for aid, name, *_ in hero.action_list if aid == action_id)


def is_need_target(hero: BasicHero, action_id: int) -> bool:
    return next(need_target for aid, _, _, need_target, *_ in hero.action_list if aid == action_id)


def serialize_stack(stack):
    return "; ".join(
        f"({a.attack_value}, {a.source.name}, {a.attack_type}, {a.damage_type})"
        for a in stack
    ) or ""


def clear_marked_actions(heroes):
    for hero in heroes:
        for stack in (hero.attack_stack, hero.defense_stack):
            for action in stack:
                if getattr(action, "marked_for_clear", False):
                    action.attack_value = 0
                    action.marked_for_clear = False


def simulate_round(p1_action_id: int, p2_action_id: int):
    p1 = BasicHero("玩家1", hp=5, hp_ceiling=5, mp=12, mp_ceiling=12)
    p2 = BasicHero("玩家2", hp=5, hp_ceiling=5, mp=12, mp_ceiling=12)
    heroes = [p1, p2]

    p1_action = (p1_action_id, action_name(p1, p1_action_id), is_need_target(p1, p1_action_id))
    p2_action = (p2_action_id, action_name(p2, p2_action_id), is_need_target(p2, p2_action_id))

    actions = [
        (p1, p1_action_id, 1 if p1_action[2] else None, next(info[4] for info in p1.action_list if info[0] == p1_action_id)),
        (p2, p2_action_id, 0 if p2_action[2] else None, next(info[4] for info in p2.action_list if info[0] == p2_action_id)),
    ]
    actions.sort(key=lambda item: ACTION_ORDER.get(item[3], 999))

    for hero in heroes:
        hero.reset_stacks()

    for hero, action_id, target_index, _ in actions:
        target = heroes[target_index] if target_index is not None else None
        hero.apply_action(action_id, target, heroes)

    clear_marked_actions(heroes)

    for hero in heroes:
        hero.update_status()

    winner: Optional[str]
    alive = [hero for hero in heroes if not hero.is_defeated()]
    if len(alive) == 2:
        winner = None
    elif len(alive) == 1:
        winner = alive[0].name
    else:
        winner = "平局"

    return {
        "p1": p1,
        "p2": p2,
        "p1_action_id": p1_action_id,
        "p1_action_name": p1_action[1],
        "p2_action_id": p2_action_id,
        "p2_action_name": p2_action[1],
        "winner": winner,
    }


def build_workbook(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    headers = [
        "p1_action_id",
        "p1_action_name",
        "p2_action_id",
        "p2_action_name",
        "winner",
        "p1_hp",
        "p1_hp_ceiling",
        "p1_mp",
        "p1_shield",
        "p1_revival_armor",
        "p2_hp",
        "p2_hp_ceiling",
        "p2_mp",
        "p2_shield",
        "p2_revival_armor",
        "p1_attack_stack",
        "p1_defense_stack",
        "p2_attack_stack",
        "p2_defense_stack",
    ]
    ws.append(headers)

    for row in results:
        p1 = row["p1"]
        p2 = row["p2"]
        ws.append([
            row["p1_action_id"],
            row["p1_action_name"],
            row["p2_action_id"],
            row["p2_action_name"],
            row["winner"],
            p1.hp,
            p1.hp_ceiling,
            p1.mp,
            p1.shield,
            p1.revival_armor,
            p2.hp,
            p2.hp_ceiling,
            p2.mp,
            p2.shield,
            p2.revival_armor,
            serialize_stack(p1.attack_stack),
            serialize_stack(p1.defense_stack),
            serialize_stack(p2.attack_stack),
            serialize_stack(p2.defense_stack),
        ])

    legend = wb.create_sheet("actions")
    legend.append(["action_id", "name", "cost", "need_target", "action_type", "attack_value", "damage_type"])
    sample_hero = BasicHero("样例")
    for action in sample_hero.action_list:
        legend.append(list(action))

    summary = wb.create_sheet("summary")
    summary.append(["p1_action_id", "p2_action_id", "winner", "p1_hp", "p2_hp"])
    for row in results:
        p1 = row["p1"]
        p2 = row["p2"]
        summary.append([row["p1_action_id"], row["p2_action_id"], row["winner"], p1.hp, p2.hp])

    wb.save(OUTPUT_FILE)


def main():
    results = []
    sample_hero = BasicHero("样例")
    action_ids = [action[0] for action in sample_hero.action_list]

    for p1_action_id in action_ids:
        for p2_action_id in action_ids:
            results.append(simulate_round(p1_action_id, p2_action_id))

    build_workbook(results)
    print(f"已生成：{OUTPUT_FILE}")
    print(f"共写入 {len(results)} 组组合")


if __name__ == "__main__":
    main()
